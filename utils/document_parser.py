import openpyxl
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    TextLoader,
)
from langchain_core.documents import Document

def parse_pdf(file_path):
    loader = PyPDFLoader(file_path)
    return loader.load()

def parse_docx(file_path):
    loader = Docx2txtLoader(file_path)
    return loader.load()

def parse_txt(file_path):
    loader = TextLoader(file_path, encoding="utf-8")
    return loader.load()

def parse_excel(file_path):
    workbook = openpyxl.load_workbook(file_path)
    text_content = []
    
    for sheet in workbook.sheetnames:
        ws = workbook[sheet]
        sheet_text = f"Sheet: {sheet}\n"
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join([str(cell) for cell in row if cell is not None])
            if row_text:
                sheet_text += row_text + "\n"
        text_content.append(sheet_text)
    
    return [Document(page_content="\n".join(text_content), metadata={"source": str(file_path)})]

def parse_document(file_path):
    ext = file_path.suffix.lower()
    parsers = {
        ".pdf": parse_pdf,
        ".docx": parse_docx,
        ".doc": parse_docx,
        ".txt": parse_txt,
        ".xlsx": parse_excel,
    }
    
    parser = parsers.get(ext)
    if not parser:
        raise ValueError(f"不支持的文件格式: {ext}")
    
    return parser(file_path)

def split_documents(documents, chunk_size=500, chunk_overlap=100):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
    )
    return splitter.split_documents(documents)
